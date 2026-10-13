from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from django.template.loader import get_template, render_to_string
from xhtml2pdf import pisa
from weasyprint import HTML
from Usuario.models import Clientes
from Proveedores.models import Proveedores
from Inventario.models import Productos
from Ventas.models import Factura
from django.contrib.auth.models import User  # Si usas usuario
from django.utils import timezone
import openpyxl
from datetime import datetime
from typing import Any
from django.templatetags.static import static
from decimal import Decimal
import pprint

# Create your views here.

def reporte_clientes_pdf(request):
    clientes = Clientes.objects.all()
    file_creation_date = datetime.now()
    report_number = f"{file_creation_date.strftime('%Y%m%d%H%M%S')}"
    html_string = render_to_string('reportes/reporte_clientes.html', {
        'clientes': clientes,
        'file_creation_date': file_creation_date,
        'report_number': report_number,
    })
    pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_clientes.pdf"'
    return response

def reporte_clientes_excel(request):
    clientes = Clientes.objects.all()
    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = 'Clientes'
    
    # Encabezados
    ws.append(['Nombre', 'Apellido', 'Teléfono', 'Email', 'Fecha de Registro'])
    
    # Datos
    for cliente in clientes:
        ws.append([cliente.nombre, cliente.apellido, cliente.telefono, cliente.email, cliente.fecha_registro])
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="clientes.xlsx"'
    wb.save(response)
    return response

def reporte_proveedores_pdf(request):
    try:
        proveedores = Proveedores.objects.all()
        template_path = 'reportes/reporte_proveedores.html'
        logo_url = request.build_absolute_uri(static('img/logo.png'))

        context = {
            'proveedores': proveedores,
            'file_creation_date': datetime.now(),
            'report_number': f"{datetime.now().strftime('%Y%m%d%H%M%S')}",  # <-- aquí el cambio
            'logo_url': logo_url,
        }

        template = get_template(template_path)
        html = template.render(context)
        try:
            pdf_file = HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf()
            response = HttpResponse(pdf_file, content_type='application/pdf')
            response['Content-Disposition'] = 'attachment; filename="reporte_proveedores.pdf"'
            return response
        except Exception:
            import traceback
            return HttpResponse(traceback.format_exc(), status=500, content_type='text/plain')
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        # Devuelve la traza completa para depuración local
        return HttpResponse(tb, status=500, content_type='text/plain')

def reporte_proveedores_excel(request):
    proveedores = Proveedores.objects.all()
    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = 'Proveedores'
    
    # Encabezados
    ws.append(['Nombre', 'Teléfono', 'Email', 'Dirección'])

    # Datos
    for proveedor in proveedores:
        ws.append([
            proveedor.nombre,
            getattr(proveedor, 'telefono', getattr(proveedor, 'Telefono', '')),
            getattr(proveedor, 'email', getattr(proveedor, 'Email', '')),
            getattr(proveedor, 'direccion', getattr(proveedor, 'Direccion', '')),
        ])
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="proveedores.xlsx"'
    wb.save(response)
    return response

def reporte_productos_pdf(request):
    try:
        productos = Productos.objects.all()
        template_path = 'reportes/reporte_productos.html'
        logo_url = request.build_absolute_uri(static('img/logo.png'))  # Agrega esto

        # Calcular totales de forma segura (evitar operaciones entre str e int)
        total_stock = 0
        valor_inventario = Decimal('0.00')
        for p in productos:
            # Cantidad: asegurar int
            qty = 0
            try:
                qty = int(getattr(p, 'Cantidad', 0) or 0)
            except Exception:
                try:
                    qty = int(str(getattr(p, 'Cantidad', 0)).strip())
                except Exception:
                    qty = 0

            # Precio: asegurar Decimal
            price = Decimal('0.00')
            try:
                price = Decimal(str(getattr(p, 'Precio', 0) or 0))
            except Exception:
                try:
                    price = Decimal(str(getattr(p, 'Precio', 0)).strip())
                except Exception:
                    price = Decimal('0.00')

            total_stock += qty
            try:
                valor_inventario += price * Decimal(qty)
            except Exception:
                # Forzar cálculo seguro: convertir qty y price
                try:
                    valor_inventario += Decimal(str(price)) * Decimal(str(qty))
                except Exception:
                    # ignora este item en el total si no se puede convertir
                    pass

        context = {
            'productos': productos,
            'file_creation_date': datetime.now(),
            'report_number': int(datetime.now().timestamp()),
            'logo_url': logo_url,
            'total_stock': total_stock,
            'valor_inventario': f"{valor_inventario:.2f}",
        }

        template = get_template(template_path)
        html = template.render(context)
        # Usar WeasyPrint en vez de xhtml2pdf/reportlab para manejar mejor tablas y CSS
        try:
            pdf_file = HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf()
            response = HttpResponse(pdf_file, content_type='application/pdf')
            response['Content-Disposition'] = 'attachment; filename="reporte_productos.pdf"'
            return response
        except Exception:
            # si WeasyPrint falla, devolver traza para depuración
            import traceback
            return HttpResponse(traceback.format_exc(), status=500, content_type='text/plain')
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        return HttpResponse(tb, status=500, content_type='text/plain')

def reporte_productos_excel(request):
    productos = Productos.objects.all()
    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = 'Productos'
    
    # Encabezados
    ws.append(['Código', 'Descripción', 'Cantidad', 'Precio', 'Fecha de creación'])

    # Datos
    for producto in productos:
        ws.append([producto.Codigo, producto.Descripcion, producto.Cantidad, producto.Precio, producto.Fecha_creacion])
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="productos.xlsx"'
    wb.save(response)
    return response

def reporte_facturas_pdf(request):
    facturas = Factura.objects.select_related('cliente').order_by('-fecha')
    response_data = {
        'facturas': facturas,
        'file_creation_date': datetime.now(),
        'report_number': f"{datetime.now().strftime('%Y%m%d%H%M%S')}"
    }
    html_string = render_to_string('reportes/reporte_facturas.html', response_data)
    pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_facturas.pdf"'
    return response


def reporte_facturas_excel(request):
    facturas = Factura.objects.select_related('cliente').order_by('-fecha')
    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = 'Facturas'
    ws.append(['Número', 'Fecha', 'Cliente', 'Cédula', 'Estado', 'Total'])
    for factura in facturas:
        cliente = getattr(factura, 'cliente', None)
        nombre = f"{getattr(cliente, 'nombre', '')} {getattr(cliente, 'apellido', '')}".strip()
        fecha = factura.fecha
        if timezone.is_aware(fecha):
            fecha = timezone.localtime(fecha).replace(tzinfo=None)
        ws.append([
            factura.numero,
            fecha,
            nombre,
            getattr(cliente, 'cedula', ''),
            factura.estado,
            getattr(factura, 'total', 0),
        ])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="facturas.xlsx"'
    wb.save(response)
    return response

def reporte_factura_pdf(request, numero):
    try:
        factura: Any = get_object_or_404(Factura, numero=numero)
        # Obtener detalles: usar related_name 'detalles' si existe, si no usar el related set por defecto
        detalles_attr = getattr(factura, 'detalles', None)
        if detalles_attr is not None:
            detalles_qs = detalles_attr.all()
        else:
            detalles_qs = getattr(factura, 'detallefactura_set', [])

        def clean_text(text):
            if not text:
                return "-"
            return str(text).replace('\n', ' ').replace('\r', ' ').replace('"', "'").replace("’", "'").replace("“", '"').replace("”", '"').strip()

        detalles_data = []
        iva_total = Decimal('0.00')
        descuento_total = Decimal('0.00')
        base_imponible_total = Decimal('0.00')
        total_calculado = Decimal('0.00')
        impuestos_por_alicuota = {}

        for d in detalles_qs:
            try:
                precio = Decimal(str(getattr(d, 'precio', 0) or 0))
            except Exception:
                precio = Decimal('0.00')
            try:
                cantidad = int(getattr(d, 'cantidad', 0) or 0)
            except Exception:
                cantidad = 0
            try:
                iva = Decimal(str(getattr(d, 'iva', 0) or 0))
            except Exception:
                iva = Decimal('0.00')
            try:
                descuento = Decimal(str(getattr(d, 'descuento', 0) or 0))
            except Exception:
                descuento = Decimal('0.00')

            line_base = Decimal(str(getattr(d, 'base_imponible', 0) or 0))
            iva_line = Decimal(str(getattr(d, 'monto_iva', 0) or 0))
            descuento_line = Decimal(str(getattr(d, 'importe_descuento', 0) or 0))
            line_subtotal = Decimal(str(getattr(d, 'total_linea', 0) or 0))

            iva_total += iva_line
            descuento_total += descuento_line
            base_imponible_total += line_base
            total_calculado += line_subtotal
            grupo_iva = impuestos_por_alicuota.setdefault(iva, {
                'alicuota': iva,
                'base_imponible': Decimal('0.00'),
                'monto_iva': Decimal('0.00'),
            })
            grupo_iva['base_imponible'] += line_base
            grupo_iva['monto_iva'] += iva_line

            detalles_data.append({
                'Codigo': getattr(d.producto, 'Codigo', ''),
                'Nombre': clean_text(getattr(d.producto, 'Descripcion', '')),
                'precio': float(precio),
                'cantidad': cantidad,
                'iva': float(iva),
                'descuento': float(descuento),
                'importe_descuento': float(descuento_line),
                'base_imponible': float(line_base),
                'monto_iva': float(iva_line),
                'subtotal': float(line_subtotal),
            })

        total = Decimal(str(getattr(factura, 'total', 0) or 0))
        resumen_impuestos = [
            impuestos_por_alicuota[alicuota]
            for alicuota in sorted(impuestos_por_alicuota)
        ]

        logo_url = request.build_absolute_uri(static('img/logo.png'))
        context = {
            'factura': factura,
            'detalles': detalles_data,
            'iva_total': float(iva_total),
            'descuento_total': float(descuento_total),
            'base_imponible': float(base_imponible_total),
            'monto_iva': float(iva_total),
            'resumen_impuestos': resumen_impuestos,
            'total_calculado': float(total_calculado),
            'total_no_concilia': abs(total - total_calculado) > Decimal('0.01'),
            'razon_social_cliente': f'{factura.cliente.nombre} {factura.cliente.apellido}'.strip(),
            'documento_cliente': factura.cliente.cedula or '',
            'domicilio_cliente': factura.cliente.direccion or '',
            'total': float(total),
            'now': datetime.now(),
            'usuario': request.user,
            'logo_url': logo_url,
        }
        html = render_to_string('reportes/reporte_facturacion.html', context)
        try:
            pdf_file = HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf()
            response = HttpResponse(pdf_file, content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="factura_{factura.numero}.pdf"'
            return response
        except Exception:
            import traceback
            return HttpResponse(traceback.format_exc(), status=500, content_type='text/plain')
    except Exception as e:
        import traceback
        tb = traceback.format_exc()
        return HttpResponse(tb, status=500, content_type='text/plain')